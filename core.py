from __future__ import annotations

import csv
import io
import re
from collections import Counter
from dataclasses import dataclass
from datetime import date, datetime
from typing import Callable, Iterable
from urllib.parse import urlparse

import requests
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from requests.adapters import HTTPAdapter
from urllib3.util.retry import Retry


SALEBOT_TOKEN_RE = re.compile(r"^[A-Za-z0-9_-]+$")
DATE_FORMATS = ("%d.%m.%Y", "%Y-%m-%d")
TIME_FORMATS = ("%H:%M", "%H:%M:%S")
REQUIRED_COLUMNS = ("client_id", "current_date", "current_time")


@dataclass(frozen=True)
class Source:
    name: str
    url: str


@dataclass
class Result:
    workbook: bytes
    source_rows: int
    result_rows: int
    removed_duplicates: int
    excluded_rows: int
    years: Counter
    source_counts: list[tuple[str, int]]
    columns: list[str]


def parse_sources(text: str) -> list[Source]:
    sources: list[Source] = []
    seen_urls: set[str] = set()
    for line_number, raw_line in enumerate(text.splitlines(), start=1):
        line = raw_line.strip()
        if not line:
            continue
        if "|" not in line:
            raise ValueError(
                f"Строка {line_number}: используйте формат «Название | ссылка»."
            )
        name, raw_url = (part.strip() for part in line.split("|", 1))
        if not name:
            raise ValueError(f"Строка {line_number}: не указано название воронки.")
        url = normalize_salebot_url(raw_url)
        if url in seen_urls:
            raise ValueError(f"Строка {line_number}: эта ссылка уже добавлена.")
        seen_urls.add(url)
        sources.append(Source(name=name, url=url))
    if not sources:
        raise ValueError("Добавьте хотя бы одну таблицу SaleBot.")
    return sources


def normalize_salebot_url(raw_url: str) -> str:
    url = raw_url.strip().rstrip("/")
    parsed = urlparse(url)
    if parsed.scheme != "https" or parsed.hostname not in {"salebot.pro", "www.salebot.pro"}:
        raise ValueError("Допустимы только HTTPS-ссылки на salebot.pro.")
    parts = [part for part in parsed.path.split("/") if part]
    if len(parts) < 3 or parts[:2] != ["shared", "table"]:
        raise ValueError("Ссылка должна иметь вид https://salebot.pro/shared/table/…")
    token = parts[2]
    if not SALEBOT_TOKEN_RE.fullmatch(token):
        raise ValueError("В ссылке SaleBot найден некорректный идентификатор таблицы.")
    return f"https://salebot.pro/shared/table/{token}"


def academic_year(day: date) -> int | None:
    if date(2026, 3, 1) <= day < date(2026, 4, 1):
        return 26
    if date(2026, 4, 1) <= day < date(2027, 5, 1):
        return 27
    return None


def parse_datetime(day_value: str, time_value: str) -> tuple[date, datetime]:
    day = None
    for fmt in DATE_FORMATS:
        try:
            day = datetime.strptime(day_value.strip(), fmt).date()
            break
        except ValueError:
            continue
    if day is None:
        raise ValueError(f"Не удалось распознать дату «{day_value}».")

    clock = None
    for fmt in TIME_FORMATS:
        try:
            clock = datetime.strptime(time_value.strip(), fmt).time()
            break
        except ValueError:
            continue
    if clock is None:
        raise ValueError(f"Не удалось распознать время «{time_value}».")
    return day, datetime.combine(day, clock)


def requests_session() -> requests.Session:
    retry = Retry(
        total=3,
        connect=3,
        read=3,
        backoff_factor=1,
        status_forcelist=(429, 500, 502, 503, 504),
        allowed_methods=frozenset({"GET"}),
    )
    session = requests.Session()
    session.headers.update({"User-Agent": "TopScoreUnique/1.0"})
    session.mount("https://", HTTPAdapter(max_retries=retry))
    return session


def download_csv(session: requests.Session, source: Source) -> str:
    export_url = f"{source.url}/export.csv"
    response = session.get(export_url, timeout=(15, 180))
    response.raise_for_status()
    content_type = response.headers.get("Content-Type", "").lower()
    if "csv" not in content_type and not response.content.startswith(b"\xef\xbb\xbf"):
        raise ValueError(f"«{source.name}»: SaleBot не вернул CSV-файл.")
    return response.content.decode("utf-8-sig")


def build_unique_workbook(
    sources: Iterable[Source],
    downloader: Callable[[Source], str] | None = None,
    progress: Callable[[int, int, str], None] | None = None,
) -> Result:
    source_list = list(sources)
    session = requests_session() if downloader is None else None
    if downloader is None:
        downloader = lambda source: download_csv(session, source)  # type: ignore[arg-type]

    headers: list[str] = []
    winners: dict[tuple[str, int], tuple[datetime, int, dict[str, str]]] = {}
    source_counts: list[tuple[str, int]] = []
    years = Counter()
    source_rows = 0
    excluded_rows = 0
    global_order = 0

    try:
        for source_index, source in enumerate(source_list, start=1):
            if progress:
                progress(source_index - 1, len(source_list), f"Загружаю: {source.name}")
            text = downloader(source)
            reader = csv.DictReader(io.StringIO(text, newline=""), delimiter=";")
            own_headers = reader.fieldnames or []
            missing = [name for name in REQUIRED_COLUMNS if name not in own_headers]
            if missing:
                raise ValueError(
                    f"«{source.name}»: отсутствуют столбцы {', '.join(missing)}."
                )
            if len(own_headers) != len(set(own_headers)):
                raise ValueError(f"«{source.name}»: названия столбцов повторяются.")
            for header in own_headers:
                if header and header not in headers:
                    headers.append(header)

            count = 0
            for csv_row_number, row in enumerate(reader, start=2):
                count += 1
                source_rows += 1
                global_order += 1
                client_id = (row.get("client_id") or "").strip()
                if not client_id:
                    excluded_rows += 1
                    continue
                try:
                    day, moment = parse_datetime(
                        row.get("current_date") or "", row.get("current_time") or ""
                    )
                except ValueError as error:
                    raise ValueError(
                        f"«{source.name}», строка {csv_row_number}: {error}"
                    ) from error
                year = academic_year(day)
                if year is None:
                    excluded_rows += 1
                    continue
                complete_row = {key: value or "" for key, value in row.items() if key}
                complete_row["Воронка"] = source.name
                key = (client_id, year)
                previous = winners.get(key)
                if previous is None or (moment, global_order) < previous[:2]:
                    winners[key] = (moment, global_order, complete_row)
            source_counts.append((source.name, count))
            if progress:
                progress(source_index, len(source_list), f"Загружено: {source.name}")
    finally:
        if session is not None:
            session.close()

    selected = sorted(winners.items(), key=lambda item: item[1][1])
    for (_, year), _ in selected:
        years[year] += 1
    output_headers = ["Воронка", *headers]
    workbook = create_xlsx(output_headers, (item[1][2] for item in selected))
    return Result(
        workbook=workbook,
        source_rows=source_rows,
        result_rows=len(selected),
        removed_duplicates=source_rows - excluded_rows - len(selected),
        excluded_rows=excluded_rows,
        years=years,
        source_counts=source_counts,
        columns=output_headers,
    )


def create_xlsx(headers: list[str], rows: Iterable[dict[str, str]]) -> bytes:
    output = io.BytesIO()
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("Уникальные подписчики")
    sheet.freeze_panes = "A2"
    for index, header in enumerate(headers, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = min(
            max(len(header) + 3, 16), 32
        )

    header_cells = []
    for header in headers:
        cell = WriteOnlyCell(sheet, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(fill_type="solid", fgColor="226A5A")
        header_cells.append(cell)
    sheet.append(header_cells)

    row_count = 0
    for row_count, row in enumerate(rows, start=1):
        cells = []
        for header in headers:
            cell = WriteOnlyCell(sheet, value=row.get(header, ""))
            cell.data_type = "s"
            cells.append(cell)
        sheet.append(cells)
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{row_count + 1}"
    workbook.save(output)
    return output.getvalue()

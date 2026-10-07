import csv
import io
from openpyxl import load_workbook

from core import Source, academic_year, build_unique_workbook, parse_sources


def make_csv(rows):
    output = io.StringIO(newline="")
    writer = csv.DictWriter(
        output,
        fieldnames=["ID", "client_id", "current_date", "current_time", "tag"],
        delimiter=";",
    )
    writer.writeheader()
    writer.writerows(rows)
    return output.getvalue()


def test_sources_and_academic_year():
    sources = parse_sources(
        "Первая | https://salebot.pro/shared/table/abc-123_X\n"
        "Вторая | https://www.salebot.pro/shared/table/second/"
    )
    assert [source.name for source in sources] == ["Первая", "Вторая"]
    assert sources[1].url == "https://salebot.pro/shared/table/second"


def test_keeps_earliest_per_client_and_year():
    source_a = Source("A", "https://salebot.pro/shared/table/a")
    source_b = Source("B", "https://salebot.pro/shared/table/b")
    data = {
        source_a: make_csv(
            [
                {"ID": "1", "client_id": "x", "current_date": "05.03.2026", "current_time": "10:00", "tag": "later"},
                {"ID": "2", "client_id": "x", "current_date": "01.04.2026", "current_time": "12:00", "tag": "year27"},
            ]
        ),
        source_b: make_csv(
            [
                {"ID": "3", "client_id": "x", "current_date": "05.03.2026", "current_time": "09:00", "tag": "earlier"},
                {"ID": "4", "client_id": "y", "current_date": "06.03.2026", "current_time": "09:00", "tag": "other"},
            ]
        ),
    }
    result = build_unique_workbook([source_a, source_b], downloader=data.__getitem__)
    assert result.source_rows == 4
    assert result.result_rows == 3
    assert result.removed_duplicates == 1
    assert result.years == {26: 2, 27: 1}

    workbook = load_workbook(io.BytesIO(result.workbook), read_only=True, data_only=True)
    rows = list(workbook.active.iter_rows(values_only=True))
    header = rows[0]
    records = [dict(zip(header, row)) for row in rows[1:]]
    assert [record["ID"] for record in records] == ["2", "3", "4"]
    assert records[1]["Воронка"] == "B"
